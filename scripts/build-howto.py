#!/usr/bin/env python3
"""Build docs/jeff-howto.md into a .docx and a .pdf. Usage: build-howto.py OUT_DIR  (needs python-docx, reportlab)."""
import re, sys, pathlib
from docx import Document
from docx.shared import Pt, Inches
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, ListFlowable, ListItem
from reportlab.lib import colors

ROOT = pathlib.Path(__file__).resolve().parent.parent
out_dir = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else ".")
name = "Jeff Clark Artworks - How to run the site"
md = (ROOT / "docs/jeff-howto.md").read_text().splitlines()

blocks, cur = [], None
for line in md:
    if line.startswith("# "): blocks.append(("h1", line[2:])); cur = None
    elif line.startswith("## "): blocks.append(("h2", line[3:])); cur = None
    elif re.match(r"^\d+\. ", line): blocks.append(("ol", re.sub(r"^\d+\. ", "", line))); cur = blocks[-1]
    elif line.startswith("- "): blocks.append(("ul", line[2:])); cur = blocks[-1]
    elif line.startswith("   - "): blocks.append(("ul2", line[5:])); cur = blocks[-1]
    elif line.strip() == "": cur = None
    else:
        if cur and cur[0] == "p": blocks[-1] = ("p", cur[1] + " " + line.strip()); cur = blocks[-1]
        else: blocks.append(("p", line.strip())); cur = blocks[-1]

def inline_docx(par, text):
    for pt in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if pt.startswith("**"): par.add_run(pt[2:-2]).bold = True
        elif pt.startswith("`"): par.add_run(pt[1:-1]).font.name = "Consolas"
        else: par.add_run(pt)

doc = Document()
doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(11)
for sec in doc.sections: sec.left_margin = sec.right_margin = Inches(1); sec.top_margin = sec.bottom_margin = Inches(0.9)
for kind, text in blocks:
    if kind == "h1": doc.add_heading(text, 0)
    elif kind == "h2": doc.add_heading(text, 1)
    elif kind == "ol": inline_docx(doc.add_paragraph(style="List Number"), text)
    elif kind == "ul": inline_docx(doc.add_paragraph(style="List Bullet"), text)
    elif kind == "ul2": inline_docx(doc.add_paragraph(style="List Bullet 2"), text)
    else: inline_docx(doc.add_paragraph(), text)
doc.core_properties.title = name; doc.core_properties.author = "Lingua Ink Media"
doc.save(out_dir / f"{name}.docx")

def inline_pdf(t):
    t = t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    t = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", t)
    return re.sub(r"`([^`]+)`", r'<font face="Courier">\1</font>', t)
ss = getSampleStyleSheet()
H1 = ParagraphStyle("H1", parent=ss["Title"], fontName="Helvetica-Bold", fontSize=20, leading=24, spaceAfter=10, alignment=0)
H2 = ParagraphStyle("H2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=13, leading=16, spaceBefore=14, spaceAfter=6, textColor=colors.HexColor("#1F1E24"))
P = ParagraphStyle("P", parent=ss["BodyText"], fontName="Helvetica", fontSize=10.5, leading=14.5, spaceAfter=6)
story, pending, ptype = [], [], None
def flush():
    global pending, ptype
    if pending:
        story.append(ListFlowable([ListItem(Paragraph(inline_pdf(t), P), leftIndent=14) for t in pending],
                                  bulletType="1" if ptype == "ol" else "bullet", start=None if ptype == "ol" else "•", leftIndent=16, bulletFontSize=9))
        pending, ptype = [], None
for kind, text in blocks:
    if kind in ("ol", "ul", "ul2"):
        k = "ol" if kind == "ol" else "ul"
        if ptype and ptype != k: flush()
        ptype = k; pending.append(("    " if kind == "ul2" else "") + text); continue
    flush()
    story.append(Paragraph(text if kind == "h1" else (text if kind == "h2" else inline_pdf(text)), {"h1": H1, "h2": H2}.get(kind, P)))
flush()
SimpleDocTemplate(str(out_dir / f"{name}.pdf"), pagesize=letter, leftMargin=inch, rightMargin=inch, topMargin=0.9*inch, bottomMargin=0.9*inch, title=name, author="Lingua Ink Media").build(story)
print("wrote", out_dir / f"{name}.docx", "and .pdf")
